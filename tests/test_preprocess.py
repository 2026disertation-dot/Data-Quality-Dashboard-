"""
System tests for the Kaggle download and cleaning pipeline.

These tests do not require the Kaggle dataset to be present: the workbook
parsing helpers are exercised against a purpose-built fixture workbook that
reproduces the layout of the real ``Portfolio WBS`` sheets, and the transform
is exercised against long-format EVM frames.
"""

import os
import sys
import tempfile

import pandas as pd
import pytest

# Add parent directory to path for imports
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from dqd.legacy.contract import validate_data_contract
from dqd.preprocess.clean import (
    CONTRACT_COLUMNS,
    build_project_totals,
    clean_and_transform_kaggle_data,
    de_cumulativise,
    load_source,
    transform_to_contract,
)
from dqd.preprocess.download import (
    FINANCE_SHEET,
    SCHEDULE_ONLY_PREFIX,
    coerce_date,
    coerce_number,
    discover_workbooks,
    find_block_period,
    find_evm_blocks,
    find_label_row,
    find_wbs_rows,
    is_finance_workbook,
    project_code_from_filename,
    project_name_from_filename,
    read_evm_workbook,
    resolve_finance_sheet,
    resolve_dataset_directory,
)

EVM_LABELS = ["PQ", "PV", "AQ", "AC", "EV"]


def _write_evm_workbook(path, sheet_name, months, wbs_lines):
    """
    Write a workbook mimicking the real 'Portfolio WBS' sheet layout.

    Row 1 holds the static columns plus a month date, row 2 holds the
    PQ/PV/AQ/AC/EV labels, and one five-column block follows per month.
    """
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.cell(row=1, column=1, value="Portfolio WBS")
    sheet.cell(row=1, column=2, value="Description")
    sheet.cell(row=1, column=3, value="Quantity")
    sheet.cell(row=1, column=4, value="BAC")

    for month_index, (date, _values) in enumerate(months):
        first_col = 5 + month_index * len(EVM_LABELS)
        sheet.cell(row=1, column=first_col + 2, value=date.to_pydatetime())
        for offset, label in enumerate(EVM_LABELS):
            sheet.cell(row=2, column=first_col + offset, value=label)

    for line_index, code in enumerate(wbs_lines):
        row = 3 + line_index
        sheet.cell(row=row, column=1, value=code)
        sheet.cell(row=row, column=2, value="Work package %s" % code)
        sheet.cell(row=row, column=3, value=10)
        sheet.cell(row=row, column=4, value=1000.0)
        for month_index, (_date, values) in enumerate(months):
            first_col = 5 + month_index * len(EVM_LABELS)
            for offset, key in enumerate(["pq", "pv", "aq", "ac", "ev"]):
                sheet.cell(row=row, column=first_col + offset, value=values[code][key])

    workbook.save(path)


def _monthly_wbs():
    """Three months of cumulative EVM values for two WBS lines."""
    return [
        (pd.Timestamp("2011-04-01"), {
            "11": {"pq": 10, "pv": 100.0, "aq": 5, "ac": 50.0, "ev": 60.0},
            "13": {"pq": 20, "pv": 200.0, "aq": 10, "ac": 120.0, "ev": 150.0},
        }),
        (pd.Timestamp("2011-05-01"), {
            "11": {"pq": 20, "pv": 250.0, "aq": 9, "ac": 140.0, "ev": 220.0},
            "13": {"pq": 40, "pv": 450.0, "aq": 18, "ac": 260.0, "ev": 400.0},
        }),
        (pd.Timestamp("2011-06-01"), {
            "11": {"pq": 20, "pv": 300.0, "aq": 10, "ac": 160.0, "ev": 300.0},
            "13": {"pq": 40, "pv": 500.0, "aq": 20, "ac": 300.0, "ev": 500.0},
        }),
    ]


@pytest.fixture
def evm_dataset_dir():
    """A dataset directory holding one financial and one schedule-only workbook."""
    from openpyxl import Workbook

    with tempfile.TemporaryDirectory() as directory:
        _write_evm_workbook(
            os.path.join(directory, "P018 - P01 Test Project.xlsx"),
            FINANCE_SHEET,
            months=_monthly_wbs(),
            wbs_lines=["11", "13"],
        )
        schedule = Workbook()
        schedule.active.title = "P01-PMB"
        schedule.active.cell(row=1, column=1, value="TASK")
        schedule.save(
            os.path.join(directory, SCHEDULE_ONLY_PREFIX + "P018 - P01 - Test.xlsx")
        )
        yield directory


@pytest.fixture
def long_evm_frame():
    """Long-format cumulative EVM data for a single project over three months."""
    rows = []
    for period_index, (date, code, pv, ac, ev) in enumerate(
        [
            (pd.Timestamp("2011-04-01"), "11", 100.0, 50.0, 60.0),
            (pd.Timestamp("2011-05-01"), "11", 250.0, 140.0, 220.0),
            (pd.Timestamp("2011-06-01"), "11", 300.0, 160.0, 300.0),
        ],
        start=1,
    ):
        rows.append({
            "project_id": "P01",
            "project_name": "Test Project",
            "source_file": "P018 - P01 Test Project.xlsx",
            "wbs_code": code,
            "wbs_description": "Work package %s" % code,
            "period_date": date,
            "period_index": period_index,
            "pv_cum": pv,
            "ac_cum": ac,
            "ev_cum": ev,
        })
    return pd.DataFrame(rows)


class TestProjectIdentifiers:
    """Project code and name extraction from workbook filenames."""

    def test_parses_project_code_ignoring_portfolio_prefix(self):
        assert project_code_from_filename("P018 - P01 Airport Carpark.xlsx") == "P01"

    def test_parses_low_numbered_project_code(self):
        assert project_code_from_filename("P038 - P17 Marina Sub-division.xlsx") == "P17"

    def test_parses_code_without_portfolio_prefix(self):
        assert project_code_from_filename("P01 Airport Carpark.xlsx") == "P01"

    def test_returns_none_when_no_code_present(self):
        assert project_code_from_filename("Unrelated Workbook.xlsx") is None

    def test_parses_project_name(self):
        assert project_name_from_filename("P018 - P01 Airport Carpark.xlsx") == "Airport Carpark"

    def test_project_name_falls_back_to_stem(self):
        assert project_name_from_filename("Unrelated Workbook.xlsx") == "Unrelated Workbook"


class TestCoercion:
    """Coercion of raw spreadsheet cells into dates and numbers."""

    def test_coerce_date_accepts_timestamp(self):
        assert coerce_date(pd.Timestamp("2011-04-01")) == pd.Timestamp("2011-04-01")

    def test_coerce_date_accepts_excel_serial(self):
        # Excel serial 40634 is 2011-04-01.
        assert coerce_date(40634) == pd.Timestamp("2011-04-01")

    def test_coerce_date_rejects_plain_number(self):
        assert coerce_date(5) is None

    def test_coerce_date_rejects_blank(self):
        assert coerce_date(None) is None
        assert coerce_date(float("nan")) is None

    def test_coerce_number_accepts_numeric_text(self):
        assert coerce_number("1234.5") == 1234.5

    def test_coerce_number_rejects_non_numeric(self):
        assert coerce_number("n/a") is None
        assert coerce_number(None) is None


class TestSheetParsing:
    """Detection of the EVM header row, monthly blocks and WBS rows."""

    def test_find_label_row(self):
        df = pd.DataFrame([[None, "PQ", "PV", "AQ", "AC", "EV"]])
        assert find_label_row(df) == 0

    def test_find_label_row_returns_none_without_labels(self):
        df = pd.DataFrame([["TASK", "DURn"], ["a", 1]])
        assert find_label_row(df) is None

    def test_find_evm_blocks_groups_monthly_columns(self):
        df = pd.DataFrame([[None] + EVM_LABELS + EVM_LABELS])
        blocks = find_evm_blocks(df, 0)
        assert len(blocks) == 2
        assert blocks[0] == {"pq": 1, "pv": 2, "aq": 3, "ac": 4, "ev": 5}
        assert blocks[1] == {"pq": 6, "pv": 7, "aq": 8, "ac": 9, "ev": 10}

    def test_find_wbs_rows_skips_headers_and_totals(self):
        df = pd.DataFrame([
            ["Portfolio WBS", "Description"],
            ["Totals, all cost codes", "SQ"],
            [11, "Survey - Construction"],
            ["", "Hire of Temporary Fencing"],
            [13, "QA Testing"],
        ])
        assert find_wbs_rows(df, 1) == [2, 4]

    def test_find_wbs_rows_accepts_numeric_codes(self):
        df = pd.DataFrame([[11.0, "Survey"], [13.0, "QA"]])
        assert find_wbs_rows(df, -1) == [0, 1]

    def test_find_block_period_reads_month_header(self):
        df = pd.DataFrame([
            [None, None, None, pd.Timestamp("2011-04-01")],
            [None] + EVM_LABELS,
        ])
        block = find_evm_blocks(df, 1)[0]
        assert find_block_period(df, block, 1) == pd.Timestamp("2011-04-01")


class TestDatasetDiscovery:
    """Workbook selection and dataset directory resolution."""

    def test_finance_workbook_is_recognised(self, evm_dataset_dir):
        path = os.path.join(evm_dataset_dir, "P018 - P01 Test Project.xlsx")
        assert is_finance_workbook(path) is True

    def test_schedule_workbook_is_rejected(self, evm_dataset_dir):
        path = os.path.join(evm_dataset_dir, SCHEDULE_ONLY_PREFIX + "P018 - P01 - Test.xlsx")
        assert is_finance_workbook(path) is False

    def test_discover_workbooks_splits_finance_from_schedule(self, evm_dataset_dir):
        finance, skipped = discover_workbooks(evm_dataset_dir)
        assert len(finance) == 1
        assert os.path.basename(finance[0]) == "P018 - P01 Test Project.xlsx"
        assert len(skipped) == 1
        assert "schedule-only" in skipped[0][1]

    def test_resolve_finance_sheet_is_case_insensitive(self):
        assert resolve_finance_sheet(["portfolio wbs", "Other"]) == "portfolio wbs"

    def test_resolve_finance_sheet_falls_back_to_wbs_name(self):
        assert resolve_finance_sheet(["Summary WBS", "Other"]) == "Summary WBS"

    def test_resolve_finance_sheet_returns_none(self):
        assert resolve_finance_sheet(["Estimate", "Resources"]) is None

    def test_resolve_dataset_directory_prefers_local_dir(self, evm_dataset_dir):
        resolved = resolve_dataset_directory(local_dir=evm_dataset_dir)
        assert resolved == os.path.abspath(evm_dataset_dir)


class TestEVMExtraction:
    """Reading a financial workbook into long-format EVM data."""

    def test_reads_one_record_per_wbs_and_month(self, evm_dataset_dir):
        path = os.path.join(evm_dataset_dir, "P018 - P01 Test Project.xlsx")
        df = read_evm_workbook(path)
        assert len(df) == 6  # 2 WBS lines x 3 months
        assert set(df["wbs_code"]) == {"11", "13"}
        assert df["project_id"].unique().tolist() == ["P01"]

    def test_preserves_cumulative_values(self, evm_dataset_dir):
        path = os.path.join(evm_dataset_dir, "P018 - P01 Test Project.xlsx")
        df = read_evm_workbook(path)
        first = df[(df["wbs_code"] == "11") & (df["period_index"] == 1)]
        assert first["pv_cum"].iloc[0] == 100.0
        assert first["ac_cum"].iloc[0] == 50.0
        assert first["ev_cum"].iloc[0] == 60.0

    def test_assigns_monthly_period_dates(self, evm_dataset_dir):
        path = os.path.join(evm_dataset_dir, "P018 - P01 Test Project.xlsx")
        df = read_evm_workbook(path)
        assert sorted(df["period_date"].unique()) == [
            pd.Timestamp("2011-04-01"),
            pd.Timestamp("2011-05-01"),
            pd.Timestamp("2011-06-01"),
        ]

    def test_rejects_workbook_without_project_code(self, evm_dataset_dir):
        from openpyxl import Workbook

        path = os.path.join(evm_dataset_dir, "Unrelated Workbook.xlsx")
        workbook = Workbook()
        workbook.active.title = FINANCE_SHEET
        workbook.save(path)
        with pytest.raises(ValueError, match="no project code"):
            read_evm_workbook(path)


class TestDeCumulativise:
    """Conversion of cumulative series into period movements."""

    def test_differences_a_rising_series(self):
        periods, negative = de_cumulativise(pd.Series([100.0, 250.0, 300.0]))
        assert periods.tolist() == [100.0, 150.0, 50.0]
        assert negative == 0

    def test_handles_a_single_period(self):
        periods, negative = de_cumulativise(pd.Series([42.0]))
        assert periods.tolist() == [42.0]
        assert negative == 0

    def test_clips_and_counts_negative_movements(self):
        periods, negative = de_cumulativise(pd.Series([100.0, 80.0, 120.0]))
        assert periods.tolist() == [100.0, 0.0, 40.0]
        assert negative == 1

    def test_treats_blanks_as_zero(self):
        periods, negative = de_cumulativise(pd.Series([100.0, None, 150.0]))
        assert periods.tolist() == [100.0, 0.0, 50.0]
        assert negative == 0


class TestProjectTotals:
    """Aggregation of WBS lines up to the project level."""

    def test_sums_wbs_lines_per_project_and_month(self, long_evm_frame):
        extra = long_evm_frame.copy()
        extra["wbs_code"] = "13"
        extra[["pv_cum", "ac_cum", "ev_cum"]] = [50.0, 25.0, 30.0]
        totals = build_project_totals(pd.concat([long_evm_frame, extra]))
        assert len(totals) == 3
        first = totals.sort_values("period_date").iloc[0]
        assert first["pv_cum"] == 150.0
        assert first["ac_cum"] == 75.0
        assert first["ev_cum"] == 90.0

    def test_drops_rows_without_a_period_date(self, long_evm_frame):
        broken = long_evm_frame.copy()
        broken.loc[0, "period_date"] = pd.NaT
        totals = build_project_totals(broken)
        assert len(totals) == 2

    def test_rejects_missing_source_columns(self):
        with pytest.raises(ValueError, match="missing required column"):
            build_project_totals(pd.DataFrame({"project_id": ["P01"]}))


class TestTransformToContract:
    """Transformation of cumulative EVM data into the dashboard contract."""

    def test_returns_contract_columns_in_order(self, long_evm_frame):
        cleaned, _report = transform_to_contract(long_evm_frame)
        assert cleaned.columns.tolist() == CONTRACT_COLUMNS

    def test_produces_one_record_per_period(self, long_evm_frame):
        cleaned, _report = transform_to_contract(long_evm_frame)
        assert len(cleaned) == 3
        assert not cleaned.duplicated(["project_id", "time_period"]).any()

    def test_names_periods_month_by_number(self, long_evm_frame):
        cleaned, _report = transform_to_contract(long_evm_frame)
        assert cleaned["time_period"].tolist() == ["Month-1", "Month-2", "Month-3"]

    def test_de_cumulativises_into_period_budget_and_cost(self, long_evm_frame):
        cleaned, _report = transform_to_contract(long_evm_frame)
        assert cleaned["pmb_budget"].tolist() == [100.0, 150.0, 50.0]
        assert cleaned["actual_cost"].tolist() == [50.0, 90.0, 20.0]
        assert cleaned["revenue_claimed"].tolist() == [60.0, 160.0, 80.0]

    def test_progress_is_share_of_baseline(self, long_evm_frame):
        cleaned, report = transform_to_contract(long_evm_frame)
        # The baseline is the highest planned value reached (300.0).
        assert report.baselines["P01"] == 300.0
        assert cleaned["progress_pct"].tolist() == [20.0, 73.33, 100.0]

    def test_period_budgets_sum_to_the_baseline(self, long_evm_frame):
        cleaned, report = transform_to_contract(long_evm_frame)
        total = cleaned["pmb_budget"].sum()
        assert total == pytest.approx(report.baselines["P01"], abs=0.01)

    def test_baseline_start_is_first_reporting_month(self, long_evm_frame):
        cleaned, _report = transform_to_contract(long_evm_frame)
        assert (cleaned["baseline_start_date"] == pd.Timestamp("2011-04-01")).all()

    def test_actual_date_is_never_before_baseline_start(self, long_evm_frame):
        cleaned, _report = transform_to_contract(long_evm_frame)
        assert (cleaned["actual_date"] >= cleaned["baseline_start_date"]).all()

    def test_output_satisfies_the_data_contract(self, long_evm_frame):
        cleaned, _report = transform_to_contract(long_evm_frame)
        is_valid, _validated, errors = validate_data_contract(cleaned)
        assert is_valid, errors

    def test_money_columns_are_never_negative(self, long_evm_frame):
        revised = long_evm_frame.copy()
        revised.loc[2, "pv_cum"] = 10.0  # source history revised downwards
        cleaned, report = transform_to_contract(revised)
        assert (cleaned["pmb_budget"] >= 0).all()
        assert (cleaned["actual_cost"] >= 0).all()
        assert report.negative_diffs >= 1

    def test_progress_is_clipped_to_the_contract_range(self, long_evm_frame):
        overrun = long_evm_frame.copy()
        overrun.loc[2, "ev_cum"] = 900.0  # earned value beyond the baseline
        cleaned, report = transform_to_contract(overrun)
        assert cleaned["progress_pct"].max() <= 100.0
        assert report.progress_clipped >= 1

    def test_orders_months_by_date_not_sheet_order(self, long_evm_frame):
        shuffled = long_evm_frame.iloc[::-1].reset_index(drop=True)
        cleaned, _report = transform_to_contract(shuffled)
        assert cleaned["actual_date"].is_monotonic_increasing

    def test_keeps_projects_independent(self, long_evm_frame):
        other = long_evm_frame.copy()
        other["project_id"] = "P02"
        cleaned, report = transform_to_contract(pd.concat([long_evm_frame, other]))
        assert report.projects == 2
        # The second project restarts from its own first period rather than
        # inheriting the first project's cumulative values.
        p02 = cleaned[cleaned["project_id"] == "P02"].iloc[0]
        assert p02["pmb_budget"] == 100.0
        assert p02["time_period"] == "Month-1"

    def test_rejects_empty_source(self):
        with pytest.raises(ValueError, match="empty"):
            transform_to_contract(pd.DataFrame(columns=list(CONTRACT_COLUMNS)))


class TestCleaningPipeline:
    """The download -> clean chain, driven from a fixture dataset."""

    def test_pipeline_produces_contract_ready_csv(self, evm_dataset_dir, tmp_path):
        import dqd.preprocess.download as download_kaggle_data

        source = str(tmp_path / "kaggle_original_data.csv")
        cleaned_path = str(tmp_path / "kaggle_cleaned_data.csv")

        original = download_kaggle_data.download_and_format_dataset(
            local_dir=evm_dataset_dir, output=source
        )
        assert original is not None
        assert os.path.exists(source)

        cleaned = clean_and_transform_kaggle_data(source=source, output=cleaned_path)
        assert cleaned is not None
        assert os.path.exists(cleaned_path)

        reloaded = pd.read_csv(cleaned_path)
        assert reloaded.columns.tolist() == CONTRACT_COLUMNS
        assert len(reloaded) == 3
        is_valid, _validated, errors = validate_data_contract(reloaded)
        assert is_valid, errors

    def test_load_source_reads_an_existing_file(self, tmp_path):
        source = tmp_path / "kaggle_original_data.csv"
        source.write_text("project_id,period_date,pv_cum,ac_cum,ev_cum\nP01,2011-04-01,1,1,1\n")
        loaded = load_source(str(source))
        assert loaded is not None
        assert loaded["project_id"].tolist() == ["P01"]
